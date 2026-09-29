# -*- coding: utf-8 -*-
"""
tc_excel_writer.py
==================
Reusable CIC-template Excel writer for manual test cases.

Column layout (matches Template/CIC_Chức_năng_A.xlsx):
    A  = ID (auto-numbered formula)
    B  = Mục đích kiểm thử
    C  = Trường hợp kiểm thử
    D  = Data test
    E  = Các bước thực hiện
    F  = Kết quả mong muốn
    G  = Mức độ nghiêm trọng
    H  = Hình ảnh
    I–N = Manual Test → WEB → Lần 1 / Ngày thực hiện / Lần 2 / Ngày thực hiện / Lần 3 / Ngày thực hiện
    O  = Kết quả test (M)
    P  = Mã lỗi
    Q  = Ghi chú

Public API
----------
    section(title)          → green full-width section header row
    subsection(name)        → pink control-group header row
    tc(...)                 → test case data row
    tc_group(...)           → grouped TC rows with merged cells (col B+C)
    tc_sub(...)             → sub-group within tc_group (col C merge)
    tc_item(...)            → individual data row within tc_sub

    class FeatureSpec(...)  → holds all data for one feature sheet
    generate_excel(...)     → main entry point, writes .xlsx file

Example
-------
    from tc_excel_writer import section, subsection, tc, FeatureSpec, generate_excel

    feature = FeatureSpec(
        sheet_title  = "Thêm mới danh mục",
        feature_name = "Thêm mới danh mục",
        feature_code = "TMDC",
        precondition = "Pre-condition:\\n- Đăng nhập với jdoe/12345678@Abc",
        rows=[
            section("Giao diện"),
            subsection("Kiểm tra Textbox Mã danh mục"),
            tc("Kiểm tra bỏ trống", "Lưu khi bỏ trống → BUS-70",
               "Mã: (rỗng)", "1. Mở form\\n2. Click Lưu",
               "Lỗi inline BUS-70", "Cao"),
        ],
        kbnv_url = "",
    )

    generate_excel([feature], "output/TestCases.xlsx")
"""

import os
import re
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

# ── Internal row-type tags ─────────────────────────────────────
_SECTION    = 'section'
_SUBSECTION = 'subsection'
_TC         = 'tc'
_TC_GROUP   = 'tc_group'

# Strips any ID-like prefix before " - " (e.g. "M1_001 - ", "M1_CHAYLAI_M01 - ", "prefix_G_D01 - ")
# Column B will ONLY show "Kiểm tra [description]" — no ID codes
_RE_TC_ID  = re.compile(r'^[\w]+(?:_[\w]+)* - (?:[A-Z]+\d+ )?')  # any WORD_WORD... - [OPTIONAL_CODE ]
_RE_NHOM   = re.compile(r'^Nhóm\s+(?:\S+\s+)?[—–]\s+')           # "Nhóm A — " or "Nhóm — "
_RE_PAREN  = re.compile(r'\s*\([^)]*\)\s*$')                      # trailing " (anything)"

# Strips any group code or STT prefix from Column C (e.g. "K01 — ", "K02 — ", "M01 — ", "C01 — ", "CONF_001 — ")
_RE_TRUONG_HOP_PREFIX = re.compile(
    r'^(?:\[?(?:Nh[oó]m\s+[A-Za-z0-9_]+\s*)?[A-Z]+[0-9]+[a-z]?\]?|\[?[A-Z]+(?:_[A-Z0-9]+)+\]?|\[?\d+\]?)\s*[—–\-:\.]\s*',
    re.UNICODE
)

def _strip_tc_id(text: str) -> str:
    """Strip any 'ID - ' prefix and existing 'Kiểm tra ' prefix, then prepend 'Kiểm tra '. Column B shows only description."""
    stripped = _RE_TC_ID.sub('', text) if text else text
    # Strip any existing 'Kiểm tra' / 'kiem tra' prefix to avoid duplication
    stripped = re.sub(r'^(?:kiểm tra|kiem tra)\s+', '', stripped, flags=re.IGNORECASE).strip() if stripped else stripped
    lowered = stripped[0].lower() + stripped[1:] if stripped else stripped
    return f'Kiểm tra {lowered}' if lowered else lowered

def _strip_nhom(text: str) -> str:
    if not text:
        return text
    cleaned, n = _RE_NHOM.subn('', text)
    if n:
        cleaned = _RE_PAREN.sub('', cleaned).strip()
        cleaned = cleaned[0].lower() + cleaned[1:] if cleaned else cleaned
        return f'Kiểm tra {cleaned}'
    return text

def _strip_truong_hop(text: str) -> str:
    """Strip any group code or STT prefix (e.g. 'K01 — ', 'M01 — ', 'C01 — ', 'CONF_001 — ') from column C."""
    if not text:
        return text
    prev = None
    current = str(text).strip()
    while prev != current:
        prev = current
        current = _RE_TRUONG_HOP_PREFIX.sub('', current).strip()
    if current:
        current = current[0].upper() + current[1:]
    return current


# ══════════════════════════════════════════════════════════════
# Public row constructors
# ══════════════════════════════════════════════════════════════

def section(title: str) -> tuple:
    """
    Green full-width section header row (e.g. "Giao diện (Phần này viết các trường hợp kiểm thử cho giao diên chung và các giao diện cho các control)", "Chức năng (Phần này viết các trường hợp kiểm thử kiểm tra các ràng buộc trong cơ sở dữ liệu và cho các luồng nghiệp vụ trong tài liệu giải pháp, tích hợp với các chức năng khác)").
    Column B text spans B:F (merged). Column A is green/empty.
    """
    return (_SECTION, title)


def subsection(name: str) -> tuple:
    """
    Pink subsection header row that identifies a control group.
    (e.g. "Kiểm tra Textbox Mã danh mục").
    """
    return (_SUBSECTION, name)


def tc(
    muc_dich:   str = None,
    truong_hop: str = None,
    data:       str = None,
    buoc:       str = '',
    ket_qua:    str = '',
    priority:   str = 'Trung bình',
) -> tuple:
    """
    Test case data row.

    Parameters
    ----------
    muc_dich   : Mục đích kiểm thử (col B). May be None for continuation rows.
    truong_hop : Trường hợp kiểm thử (col C). May be None.
    data       : Data test (col D). May be None.
    buoc       : Các bước thực hiện (col E).
    ket_qua    : Kết quả mong muốn (col F).
    priority   : Mức độ nghiêm trọng (col G): "Cao" | "Trung bình" | "Thấp".
    """
    return (_TC, muc_dich, truong_hop, data, buoc, ket_qua, priority)


def tc_group(muc_dich: str, sub_groups: list) -> tuple:
    """
    Grouped test case rows with merged cells (col B + col C).

    Creates a hierarchical layout with 3 levels:
      - Level 1 (col B): ``muc_dich`` — merged vertically across all items.
      - Level 2 (col C): ``truong_hop`` in each sub-group — merged vertically.
      - Level 3 (col D): Individual data-test value per row.

    Parameters
    ----------
    muc_dich   : Level-1 group label for col B (e.g. "Kiểm tra khi nhập giá trị").
    sub_groups : list of ``tc_sub(...)`` tuples.

    Example
    -------
    ::

        tc_group("Kiểm tra khi nhập giá trị", [
            tc_sub("Kiểm tra khi nhập khoảng trắng", [
                tc_item("nhập toàn khoảng trắng",         "1. Nhập…", "Lỗi…", "Cao"),
                tc_item("nhập khoảng trắng ở đầu cuối",   "1. Nhập…", "Trim…", "Cao"),
            ]),
            tc_sub("Kiểm tra nhập maxlength", [
                tc_item("Nhập <= 50 ký tự", "1. Nhập…", "OK", "Cao"),
                tc_item("Nhập > 50 ký tự",  "1. Nhập…", "Chặn", "Cao"),
            ]),
        ])
    """
    return (_TC_GROUP, muc_dich, sub_groups)


def tc_sub(truong_hop: str, items: list) -> tuple:
    """
    Sub-group within a ``tc_group``.

    Parameters
    ----------
    truong_hop : Level-2 label for col C (e.g. "Kiểm tra khi nhập khoảng trắng").
    items      : list of ``tc_item(...)`` tuples.
    """
    return (truong_hop, items)


def tc_item(
    data:     str = '',
    buoc:     str = '',
    ket_qua:  str = '',
    priority: str = 'Trung bình',
) -> tuple:
    """
    Individual data-test row inside a ``tc_sub``.

    Parameters
    ----------
    data     : Data test value (col D).
    buoc     : Các bước thực hiện (col E).
    ket_qua  : Kết quả mong muốn (col F).
    priority : Mức độ nghiêm trọng (col G).
    """
    return (data, buoc, ket_qua, priority)


# ══════════════════════════════════════════════════════════════
# Feature specification
# ══════════════════════════════════════════════════════════════

class FeatureSpec:
    """
    All data required to build one feature sheet.

    Parameters
    ----------
    sheet_title   : Excel tab name (e.g. "Thêm mới danh mục")
    feature_name  : Full name shown in metadata row C3
    feature_code  : Short code shown in C4 (e.g. "TMDC")
    precondition  : Pre-condition text block for row 18 (yellow)
    rows          : list of section() / subsection() / tc() tuples
    kbnv_url      : Optional hyperlink text for KBNV cell F3
    """
    def __init__(
        self,
        sheet_title:   str,
        feature_name:  str,
        feature_code:  str,
        precondition:  str,
        rows:          list,
        kbnv_url:      str = '',
    ):
        self.sheet_title  = sheet_title
        self.feature_name = feature_name
        self.feature_code = feature_code
        self.precondition = precondition
        self.rows         = rows
        self.kbnv_url     = kbnv_url

    @property
    def tc_count(self) -> int:
        count = 0
        for r in self.rows:
            if r[0] == _TC:
                count += 1
            elif r[0] == _TC_GROUP:
                for sg in r[2]:          # sg = (truong_hop, items)
                    count += len(sg[1])
        return count


# ══════════════════════════════════════════════════════════════
# Style constants  (private)
# ══════════════════════════════════════════════════════════════

def _side(color: str = 'B2B2B2') -> Side:
    return Side(style='thin', color=color)


def _border() -> Border:
    s = _side()
    return Border(left=s, right=s, top=s, bottom=s)


# Fill colours (exact hex from template)
BG_HEADER     = PatternFill('solid', fgColor='FF729FCF')   # Blue  – column headers rows 15-17
BG_PRECOND    = PatternFill('solid', fgColor='FFFFFF00')   # Yellow – pre-condition row 18
BG_SECTION    = PatternFill('solid', fgColor='FFCCFFCC')   # Green  – section header + col A
BG_SUBSECTION = PatternFill('solid', fgColor='FFF4CCCC')   # Pink   – subsection header
BG_WHITE      = PatternFill('solid', fgColor='FFFFFFFF')   # White  – step/result cells
BG_TONG_GREEN = PatternFill('solid', fgColor='FF92D050')   # Tổng hợp STT/Name header
BG_TONG_BLUE  = PatternFill('solid', fgColor='FF00B0F0')   # Tổng hợp stats header


def _font(bold: bool = False, size: int = 12, color: str = '000000') -> Font:
    return Font(name='Times New Roman', bold=bold, size=size, color=color)


FONT_BOLD   = _font(bold=True,  size=12)
FONT_NORMAL = _font(size=12)
FONT_LINK   = _font(size=12, color='0000FF')

ALIGN_CC = Alignment(horizontal='center', vertical='center', wrap_text=True)
ALIGN_LC = Alignment(horizontal='left',   vertical='center', wrap_text=True)
ALIGN_LB = Alignment(horizontal='left',   vertical='center', wrap_text=True)
ALIGN_LT = Alignment(horizontal='left',   vertical='center', wrap_text=True)
ALIGN_CB = Alignment(horizontal='center', vertical='center', wrap_text=True)

# Column widths matching the original template
_COL_WIDTHS = {
    'A': 13.0,  'B': 33.63, 'C': 28.0, 'D': 28.0,
    'E': 44.25, 'F': 59.75, 'G': 12.63, 'H': 13.0,
    'I': 13.0,  'J': 13.0,  'K': 13.0,  'L': 13.0,
    'M': 13.0,  'N': 13.0,  'O': 13.0,  'P': 13.0, 'Q': 13.0,
}


# ══════════════════════════════════════════════════════════════
# Private helpers
# ══════════════════════════════════════════════════════════════

def _sc(cell, val=None, bg=None, font=None, align=None, brd: bool = True):
    """Apply value + styles to a single cell."""
    if val is not None:
        cell.value = val
    if bg:
        cell.fill = bg
    if font:
        cell.font = font
    if align:
        cell.alignment = align
    if brd:
        cell.border = _border()


def _id_formula(row: int, data_start: int, first_tc_row: int) -> str:
    """
    Auto-numbering formula for column A.
    Returns "" for section/subsection rows (F is empty).
    Returns sequential integer for TC rows.
    """
    return (
        f'=IF(F{row}="","",' \
        f'ROW()-{data_start}-COUNTBLANK($F${first_tc_row}:F{row}))'
    )


# ══════════════════════════════════════════════════════════════
# Feature sheet builder  (private)
# ══════════════════════════════════════════════════════════════

def _build_feature_sheet(wb: openpyxl.Workbook, spec: FeatureSpec):
    ws = wb.create_sheet(title=spec.sheet_title)

    # Column widths
    for col, w in _COL_WIDTHS.items():
        ws.column_dimensions[col].width = w

    # ── Row 1 – banner title ──────────────────────────────────
    ws.merge_cells('B1:F1')
    _sc(ws['B1'], val='KỊCH BẢN KIỂM THỬ CHỨC NĂNG',
        bg=BG_WHITE, font=FONT_BOLD, align=ALIGN_CC)

    # ── Rows 3-4 – metadata ───────────────────────────────────
    _sc(ws['B3'], val='Tên màn hình/Tên chức năng', bg=BG_WHITE, font=FONT_BOLD, align=ALIGN_LB)
    _sc(ws['C3'], val=spec.feature_name,             bg=BG_WHITE, font=FONT_NORMAL, align=ALIGN_LC)
    _sc(ws['E3'], val='KBNV',                         bg=BG_WHITE, font=FONT_BOLD, align=ALIGN_LB)
    _sc(ws['F3'], val=spec.kbnv_url,                  bg=BG_WHITE, font=FONT_LINK, align=ALIGN_LC)
    _sc(ws['B4'], val='Mã trường hợp kiểm thử',      bg=BG_WHITE, font=FONT_BOLD, align=ALIGN_LB)
    _sc(ws['C4'], val=spec.feature_code,              bg=BG_WHITE, font=FONT_NORMAL, align=ALIGN_LC)
    _sc(ws['E4'], val='TKCT',                          bg=BG_WHITE, font=FONT_BOLD, align=ALIGN_LB)

    # ── Row 6 – manual test label ─────────────────────────────
    _sc(ws['B6'], val='Manual test', bg=BG_WHITE, font=FONT_BOLD, align=ALIGN_LB)

    # ── Rows 7-11 – statistics ────────────────────────────────
    stats = [
        (7,  'Số trường hợp kiểm thử đạt (P)',           ''),
        (8,  'Số trường hợp kiểm thử không đạt (F)',      ''),
        (9,  'Số trường hợp kiểm thử đang xem xét (PE)', ''),
        (10, 'Số trường hợp kiểm thử chưa thực hiện',    '=C11-C7-C8-C9'),
        (11, 'Tổng số trường hợp kiểm thử',              ''),
    ]
    for r, label, formula in stats:
        _sc(ws.cell(r, 2), val=label,   bg=BG_WHITE,   font=FONT_BOLD,   align=ALIGN_LB)
        _sc(ws.cell(r, 3), val=formula, bg=BG_SECTION, font=FONT_NORMAL, align=ALIGN_CC)

    # ── Rows 15-17 – 3-level column headers ──────────────────
    header_merges = [
        ('A15:A17', 'ID'),
        ('B15:B17', 'Mục đích kiểm thử'),
        ('C15:C17', 'Trường hợp kiểm thử'),
        ('D15:D17', 'Data test'),
        ('E15:E17', 'Các bước thực hiện'),
        ('F15:F17', 'Kết quả mong muốn'),
        ('G15:G17', 'Mức độ nghiêm trọng'),
        ('H15:H17', 'Hình ảnh'),
        ('I15:N15', 'Manual Test'),
        ('O15:O17', 'Kết quả test (M)'),
        ('P15:P17', 'Mã lỗi'),
        ('Q15:Q17', 'Ghi chú'),
    ]
    for cell_range, label in header_merges:
        ws.merge_cells(cell_range)
        top_left = cell_range.split(':')[0]
        _sc(ws[top_left], val=label, bg=BG_HEADER, font=FONT_BOLD, align=ALIGN_CC)

    ws.merge_cells('I16:N16')
    _sc(ws['I16'], val='WEB', bg=BG_HEADER, font=FONT_BOLD, align=ALIGN_CC)
    for col, label in [('I', 'Lần 1'), ('J', 'Ngày thực hiện'), ('K', 'Lần 2'),
                       ('L', 'Ngày thực hiện'), ('M', 'Lần 3'), ('N', 'Ngày thực hiện')]:
        _sc(ws[f'{col}17'], val=label, bg=BG_HEADER, font=FONT_BOLD, align=ALIGN_CC)

    # ── Row 18 – pre-condition (yellow) ───────────────────────
    ws.merge_cells('B18:F18')
    _sc(ws['B18'], val=spec.precondition, bg=BG_PRECOND, font=FONT_BOLD, align=ALIGN_LB)
    ws.row_dimensions[18].height = 80

    # ── Data rows – starting at row 19 ────────────────────────
    DATA_START  = 19
    current_row = DATA_START
    first_tc_row = DATA_START  # updated below to first actual TC row
    first_tc_set = False

    for item in spec.rows:

        if item[0] == _SECTION:
            # Set all values BEFORE merging (openpyxl restriction)
            _sc(ws.cell(current_row, 1), val='',      bg=BG_SECTION, font=FONT_BOLD, align=ALIGN_LB)
            _sc(ws.cell(current_row, 2), val=item[1], bg=BG_SECTION, font=FONT_BOLD, align=ALIGN_LB)
            for c in range(3, 18):
                _sc(ws.cell(current_row, c), val='', bg=BG_SECTION, font=FONT_BOLD, align=ALIGN_LB)
            ws.merge_cells(f'B{current_row}:F{current_row}')
            ws.row_dimensions[current_row].height = 20
            if not first_tc_set:
                first_tc_row = current_row + 1
                first_tc_set = True
            current_row += 1

        elif item[0] == _SUBSECTION:
            r = current_row
            _sc(ws.cell(r, 1),
                val=_id_formula(r, DATA_START, first_tc_row),
                bg=BG_SECTION, font=FONT_NORMAL, align=ALIGN_CB)
            _sc(ws.cell(r, 2), val=_strip_nhom(item[1]), bg=BG_SUBSECTION, font=FONT_BOLD, align=ALIGN_LB)
            for c in range(3, 18):
                _sc(ws.cell(r, c), val='', bg=BG_SUBSECTION, font=FONT_BOLD, align=ALIGN_LB)
            ws.row_dimensions[r].height = 20
            current_row += 1

        elif item[0] == _TC:
            _, muc_dich, truong_hop, data_test, cac_buoc, ket_qua, muc_do = item
            r = current_row
            _sc(ws.cell(r, 1),
                val=_id_formula(r, DATA_START, first_tc_row),
                bg=BG_SECTION, font=FONT_NORMAL, align=ALIGN_CB)
            _sc(ws.cell(r, 2), val=_strip_tc_id(muc_dich), bg=None, font=FONT_NORMAL, align=ALIGN_LC)
            _sc(ws.cell(r, 3), val=_strip_truong_hop(truong_hop), bg=None, font=FONT_NORMAL, align=ALIGN_LC)
            _sc(ws.cell(r, 4), val=data_test,  bg=None,     font=FONT_NORMAL, align=ALIGN_LC)
            _sc(ws.cell(r, 5), val=cac_buoc,   bg=BG_WHITE, font=FONT_NORMAL, align=ALIGN_LT)
            _sc(ws.cell(r, 6), val=ket_qua,    bg=BG_WHITE, font=FONT_NORMAL, align=ALIGN_LC)
            _sc(ws.cell(r, 7), val=muc_do,     bg=None,     font=FONT_NORMAL, align=ALIGN_CC)
            for c in range(8, 18):
                _sc(ws.cell(r, c), val='', bg=None, font=FONT_NORMAL, align=ALIGN_CC)
            ws.row_dimensions[r].height = 60
            current_row += 1

        elif item[0] == _TC_GROUP:
            _, muc_dich_grp, sub_groups = item
            # Calculate total rows for col-B merge
            total_items = sum(len(sg[1]) for sg in sub_groups)
            group_start_row = current_row

            for sg in sub_groups:
                truong_hop_sg, items_sg = sg
                sg_start_row = current_row

                for itm in items_sg:
                    data_test, cac_buoc, ket_qua, muc_do = itm
                    r = current_row

                    # Col A — ID
                    _sc(ws.cell(r, 1),
                        val=_id_formula(r, DATA_START, first_tc_row),
                        bg=BG_SECTION, font=FONT_NORMAL, align=ALIGN_CB)

                    # Col B — muc_dich (value only on first row of group)
                    val_b = muc_dich_grp if r == group_start_row else ''
                    _sc(ws.cell(r, 2), val=val_b, bg=None, font=FONT_NORMAL, align=ALIGN_LC)

                    # Col C — truong_hop (value only on first row of sub-group)
                    val_c = _strip_truong_hop(truong_hop_sg) if r == sg_start_row else ''
                    _sc(ws.cell(r, 3), val=val_c, bg=None, font=FONT_NORMAL, align=ALIGN_LC)

                    # Cols D–G — individual per row
                    _sc(ws.cell(r, 4), val=data_test,  bg=None,     font=FONT_NORMAL, align=ALIGN_LC)
                    _sc(ws.cell(r, 5), val=cac_buoc,   bg=BG_WHITE, font=FONT_NORMAL, align=ALIGN_LT)
                    _sc(ws.cell(r, 6), val=ket_qua,    bg=BG_WHITE, font=FONT_NORMAL, align=ALIGN_LC)
                    _sc(ws.cell(r, 7), val=muc_do,     bg=None,     font=FONT_NORMAL, align=ALIGN_CC)
                    for c in range(8, 18):
                        _sc(ws.cell(r, c), val='', bg=None, font=FONT_NORMAL, align=ALIGN_CC)
                    ws.row_dimensions[r].height = 60
                    current_row += 1

                # Merge col C for this sub-group (only when > 1 row)
                if len(items_sg) > 1:
                    ws.merge_cells(f'C{sg_start_row}:C{current_row - 1}')

            # Merge col B for the entire group (only when > 1 row)
            if total_items > 1:
                ws.merge_cells(f'B{group_start_row}:B{current_row - 1}')

    # ── Update statistics formulas with dynamic row range ────────
    last_row = max(current_row - 1, 20)
    _sc(ws['C7'],  val=f'=COUNTIF($O$20:$O${last_row},"P")',   bg=BG_SECTION, font=FONT_NORMAL, align=ALIGN_CC)
    _sc(ws['C8'],  val=f'=COUNTIF($O$20:$O${last_row},"F")',   bg=BG_SECTION, font=FONT_NORMAL, align=ALIGN_CC)
    _sc(ws['C9'],  val=f'=COUNTIF($O$20:$O${last_row},"PE")',  bg=BG_SECTION, font=FONT_NORMAL, align=ALIGN_CC)
    _sc(ws['C10'], val='=C11-C7-C8-C9',                        bg=BG_SECTION, font=FONT_NORMAL, align=ALIGN_CC)
    _sc(ws['C11'], val=f'=COUNTA($F$20:$F${last_row})',        bg=BG_SECTION, font=FONT_NORMAL, align=ALIGN_CC)

    return ws


# ══════════════════════════════════════════════════════════════
# Tổng hợp sheet builder  (private)
# ══════════════════════════════════════════════════════════════

def _build_tonghop_sheet(wb: openpyxl.Workbook, features: list):
    ws = wb.create_sheet(title='Tổng hợp', index=0)

    col_widths = {
        'A': 13.0, 'B': 40.0, 'C': 14.0, 'D': 14.0, 'E': 14.0,
        'F': 14.0, 'G': 14.0, 'H': 14.0, 'I': 14.0, 'J': 14.0,
    }
    for col, w in col_widths.items():
        ws.column_dimensions[col].width = w

    # Row 1 – banner title
    ws.merge_cells('A1:J1')
    _sc(ws['A1'], val='TỔNG HỢP KẾT QUẢ KIỂM THỬ', font=FONT_BOLD, align=ALIGN_CC)

    # Rows 3-4 – two-level column headers
    ws.merge_cells('A3:A4')
    _sc(ws['A3'], val='STT', bg=BG_TONG_GREEN, font=FONT_BOLD, align=ALIGN_CC)
    ws.merge_cells('B3:B4')
    _sc(ws['B3'], val='Tên màn hình/chức năng', bg=BG_TONG_GREEN, font=FONT_BOLD, align=ALIGN_CC)
    ws.merge_cells('C3:J3')
    _sc(ws['C3'], val='Manual test', bg=BG_TONG_BLUE, font=FONT_BOLD, align=ALIGN_CC)

    sub_hdrs = [
        ('C4', 'Số trường hợp kiểm thử đạt (P)'),
        ('D4', 'Số trường hợp kiểm thử không đạt (F)'),
        ('E4', 'Số trường hợp kiểm thử đang xem xét (PE)'),
        ('F4', 'Số trường hợp kiểm thử chưa thực hiện'),
        ('G4', 'Tổng số trường hợp kiểm thử'),
        ('H4', 'Tỉ lệ trường hợp kiểm thử đạt (%P)'),
        ('I4', 'Tỉ lệ trường hợp kiểm thử không đạt (%F)'),
        ('J4', 'Tỉ lệ trường hợp kiểm thử đã thực hiện (%Cover)'),
    ]
    for addr, label in sub_hdrs:
        _sc(ws[addr], val=label, bg=BG_TONG_BLUE, font=FONT_BOLD, align=ALIGN_CC)

    # Feature data rows (start at row 5)
    for i, spec in enumerate(features, 1):
        r = i + 4
        safe = spec.sheet_title.replace("'", "''")
        _sc(ws[f'A{r}'], val=i,                                       font=FONT_NORMAL, align=ALIGN_CC)
        _sc(ws[f'B{r}'], val=spec.feature_name if spec.feature_name else f"='{safe}'!C3", font=FONT_NORMAL, align=ALIGN_LC)
        for col, ref in [('C', 'C7'), ('D', 'C8'), ('E', 'C9'), ('F', 'C10'), ('G', 'C11')]:
            _sc(ws[f'{col}{r}'], val=f"='{safe}'!{ref}",              font=FONT_NORMAL, align=ALIGN_CC)
        _sc(ws[f'H{r}'], val=f'=IF(G{r}=0,"",C{r}/G{r})',            font=FONT_BOLD, align=ALIGN_CC)
        _sc(ws[f'I{r}'], val=f'=IF(G{r}=0,"",D{r}/G{r})',            font=FONT_BOLD, align=ALIGN_CC)
        _sc(ws[f'J{r}'], val=f'=IF(G{r}=0,"",(C{r}+D{r}+E{r})/G{r})', font=FONT_BOLD, align=ALIGN_CC)
        for letter in ('H', 'I', 'J'):
            ws[f'{letter}{r}'].number_format = '0.0%'

    # Total row
    last_data = len(features) + 4
    total_r = last_data + 1
    ws.merge_cells(f'A{total_r}:B{total_r}')
    _sc(ws[f'A{total_r}'], val='Total', font=FONT_BOLD, align=ALIGN_CC)
    for col in ('C', 'D', 'E', 'F', 'G'):
        _sc(ws[f'{col}{total_r}'],
            val=f'=SUM({col}5:{col}{last_data})',
            font=FONT_BOLD, align=ALIGN_CC)
    _sc(ws[f'H{total_r}'], val=f'=IF(G{total_r}=0,"",C{total_r}/G{total_r})', font=FONT_BOLD, align=ALIGN_CC)
    _sc(ws[f'I{total_r}'], val=f'=IF(G{total_r}=0,"",D{total_r}/G{total_r})', font=FONT_BOLD, align=ALIGN_CC)
    _sc(ws[f'J{total_r}'], val=f'=IF(G{total_r}=0,"",(C{total_r}+D{total_r}+E{total_r})/G{total_r})', font=FONT_BOLD, align=ALIGN_CC)
    for letter in ('H', 'I', 'J'):
        ws[f'{letter}{total_r}'].number_format = '0.0%'

    return ws


# ══════════════════════════════════════════════════════════════
# Public entry point
# ══════════════════════════════════════════════════════════════

def generate_excel(features: list, output_path: str) -> None:
    """
    Generate a CIC-template Excel file.

    Parameters
    ----------
    features    : list[FeatureSpec]  — one FeatureSpec per feature sheet
    output_path : str                — destination .xlsx path (created if missing)
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove the default blank sheet

    for spec in features:
        _build_feature_sheet(wb, spec)

    _build_tonghop_sheet(wb, features)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    try:
        wb.save(output_path)
    except PermissionError:
        import time
        new_path = output_path.replace(".xlsx", f"_{int(time.time())}.xlsx")
        wb.save(new_path)
        print(f"\n[WARNING] File {output_path} is currently locked by another program. Saved to new file: {new_path}")
        output_path = new_path

    total_tc = sum(spec.tc_count for spec in features)
    # Use sys.stdout.buffer for safe UTF-8 output on Windows consoles
    import sys
    out = sys.stdout.buffer if hasattr(sys.stdout, 'buffer') else None
    msg = (
        f'[OK] Saved  : {output_path}\n'
        f'     Sheets  : {len(features)} feature(s) + Tong hop\n'
        f'     Total TC: {total_tc}\n'
    )
    if out:
        out.write(msg.encode('utf-8'))
    else:
        print(msg, end='')
